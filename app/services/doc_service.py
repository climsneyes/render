import re
from datetime import datetime
from typing import List, Dict, Any, Optional
from docx import Document
from docx.shared import Mm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH


def create_word_document(query: str, results: List[Dict[str, Any]]) -> Document:
    """조례 검색 결과로 3단 비교 Word 문서 생성"""
    doc = Document()
    section = doc.sections[-1]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(420)
    section.page_height = Mm(297)

    title = doc.add_heading('조례 검색 결과', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f'검색어: {query}')
    doc.add_paragraph(f'총 {len(results)}건의 조례가 검색되었습니다.\n')

    # 조례를 3개씩 그룹화하여 3단 비교표 형태로 생성
    for i in range(0, len(results), 3):
        current_laws = results[i:i+3]
        while len(current_laws) < 3:
            current_laws.append({'name': '', 'content': [], 'metro': ''})

        table = doc.add_table(rows=1, cols=3)
        table.style = 'Table Grid'
        table.autofit = True

        for idx, law in enumerate(current_laws):
            cell = table.cell(0, idx)
            paragraph = cell.paragraphs[0]
            if law['name']:
                run = paragraph.add_run(f"{law['metro']}\n{law['name']}\n\n")
                run.bold = True
                run.font.color.rgb = RGBColor(255, 0, 0)

                if law['content']:
                    content_text = '\n\n'.join(law['content'])
                    paragraph.add_run(content_text)
                else:
                    paragraph.add_run('(조문 없음)')

        if i + 3 < len(results):
            doc.add_page_break()

    return doc


def parse_table_from_text(text_content: str) -> List[Dict[str, Any]]:
    """텍스트에서 표 형태의 내용을 파싱하여 Word 표 데이터로 변환"""
    tables_data = []
    lines = text_content.split('\n')
    current_table = None

    for line in lines:
        line = line.strip()
        if not line:
            if current_table is not None:
                if current_table.get('headers') or current_table.get('rows'):
                    tables_data.append(current_table)
                current_table = None
            continue

        if '|' in line:
            cells = [cell.strip() for cell in line.split('|')]
            if cells and cells[0] == '':
                cells.pop(0)
            if cells and cells[-1] == '':
                cells.pop()

            is_divider = len(cells) > 0 and all(cell.replace('-', '').replace(':', '').strip() == '' for cell in cells)

            if len(cells) >= 2:
                if current_table is None:
                    if not is_divider:
                        current_table = {'headers': cells, 'rows': []}
                else:
                    if not is_divider:
                        current_table['rows'].append(cells)
        else:
            if current_table is not None:
                if current_table.get('headers') or current_table.get('rows'):
                    tables_data.append(current_table)
                current_table = None

    if current_table is not None:
        if current_table.get('headers') or current_table.get('rows'):
            tables_data.append(current_table)

    return tables_data


def add_table_to_doc(doc: Document, table_data: Dict[str, Any]):
    """Word 문서에 표 추가"""
    headers = table_data.get('headers', [])
    rows = table_data.get('rows', [])
    if not headers and not rows:
        return

    if not headers and rows:
        headers = rows[0]
        rows = rows[1:]

    max_cols = max(len(headers), max((len(r) for r in rows), default=0))
    if max_cols == 0:
        return

    row_count = 1 + len(rows)
    table = doc.add_table(rows=row_count, cols=max_cols)
    table.style = 'Table Grid'
    table.autofit = True

    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        if i < len(header_cells):
            header_cells[i].text = header
            paragraph = header_cells[i].paragraphs[0]
            run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
            run.bold = True

    for row_idx, row_data in enumerate(rows):
        if row_idx + 1 < len(table.rows):
            cells = table.rows[row_idx + 1].cells
            for col_idx, cell_data in enumerate(row_data):
                if col_idx < len(cells):
                    cells[col_idx].text = cell_data


def add_markdown_content_to_doc(doc: Document, content: str):
    """마크다운 본문을 구조화하여 Word 문서에 삽입"""
    if not content:
        return

    lines = content.split('\n')
    i = 0
    total_lines = len(lines)

    while i < total_lines:
        line = lines[i].strip()
        if not line:
            i += 1
            continue

        if '|' in line and len([c for c in line.split('|') if c.strip()]) >= 2:
            table_lines = []
            while i < total_lines and '|' in lines[i]:
                t_line = lines[i].strip()
                if t_line:
                    table_lines.append(t_line)
                i += 1

            if table_lines:
                parsed_tables = parse_table_from_text('\n'.join(table_lines))
                if parsed_tables and any(t.get('headers') for t in parsed_tables):
                    for tbl in parsed_tables:
                        add_table_to_doc(doc, tbl)
                        doc.add_paragraph("")
                else:
                    for t_line in table_lines:
                        doc.add_paragraph(t_line)
            continue

        if line.startswith(('# ', '## ', '### ', '#### ')):
            level = min(line.count('#', 0, 5), 4)
            clean_title = re.sub(r'^#+\s*', '', line).strip()
            doc.add_heading(clean_title, level=min(level + 1, 4))
            i += 1
            continue

        # 일반 단락
        p = doc.add_paragraph()
        # 간단한 볼드 파싱 (**볼드**)
        parts = re.split(r'(\*\*.*?\*\*)', line)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                run = p.add_run(part[2:-2])
                run.bold = True
            else:
                p.add_run(part)
        i += 1


def create_comparison_document(
    pdf_text: str,
    search_results: List[Dict[str, Any]],
    analysis_text: str,
    superior_laws_content: Optional[List[Dict[str, Any]]] = None,
    relevant_guidelines: Optional[List[Dict[str, Any]]] = None
) -> Document:
    """조례 비교 분석 결과 보고서 Word 문서 생성"""
    doc = Document()
    section = doc.sections[-1]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(420)
    section.page_height = Mm(297)

    title = doc.add_heading('조례 비교 분석 결과 보고서', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f'분석 일시: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}\n')

    if superior_laws_content:
        doc.add_heading('검토된 상위법령', level=2)
        for law_group in superior_laws_content:
            base_name = law_group['base_name']
            doc.add_paragraph(f"◆ {base_name}")
            if 'combined_content' in law_group:
                content_length = len(law_group['combined_content'])
                doc.add_paragraph(f"  • 본문 {content_length:,}자")
            else:
                for law_type, law_info in law_group['laws'].items():
                    if law_info and 'articles' in law_info:
                        type_name = "법률" if law_type == 'law' else ("시행령" if law_type == 'decree' else "시행규칙")
                        doc.add_paragraph(f"  • {law_info['law_name']} ({type_name}) - {len(law_info['articles'])}개 조문")
                combined_articles = law_group.get('combined_articles', [])
                doc.add_paragraph(f"  총 {len(combined_articles)}개 조문 통합 검토")
            doc.add_paragraph("")
        doc.add_paragraph("")

    if relevant_guidelines:
        doc.add_heading('활용된 자치법규 참고자료', level=2)
        source_groups = {}
        for guideline in relevant_guidelines:
            source_store = guideline.get('source_store', '알 수 없는 자료')
            if source_store not in source_groups:
                source_groups[source_store] = []
            source_groups[source_store].append(guideline)

        for source_store, guidelines in source_groups.items():
            doc.add_paragraph(f"◆ {source_store} ({len(guidelines)}개 내용)")
            for i, guideline in enumerate(guidelines):
                similarity_score = guideline.get('similarity', 1 - guideline.get('distance', 0))
                doc.add_paragraph(f"   • 내용 {i+1} (유사도: {similarity_score:.3f})")
        doc.add_paragraph("")

    doc.add_heading('📋 AI 종합 분석 보고서', level=2)
    if analysis_text and analysis_text.strip():
        add_markdown_content_to_doc(doc, analysis_text)
    else:
        doc.add_paragraph('분석 내용이 비어있습니다.')

    return doc

